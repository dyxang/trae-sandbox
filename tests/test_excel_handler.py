import os
import tempfile
import pytest
from openpyxl import Workbook, load_workbook
from src.excel_handler import ExcelHandler


INPUT_FILE = "/workspace/templates/test_input.xlsx"


def test_detect_structure():
    handler = ExcelHandler(INPUT_FILE)
    structure = handler.detect_structure()
    assert structure["sheet_name"] == "工程量清单"
    assert structure["header_row"] == 1
    assert structure["data_start_row"] == 2
    assert structure["col_code"] == 2
    assert structure["col_name"] == 3
    assert structure["col_price"] == 7
    assert structure["total_rows"] > 0


def test_read_data():
    handler = ExcelHandler(INPUT_FILE)
    structure = handler.detect_structure()
    data = handler.read_data(structure)
    assert len(data) == structure["total_rows"]
    for item in data:
        assert "row" in item
        assert "code" in item
        assert "name" in item
        assert "price" in item
    codes = [d["code"] for d in data]
    assert "010101001001" in codes
    names = [d["name"] for d in data]
    assert "土方开挖" in names


def test_write_results():
    handler = ExcelHandler(INPUT_FILE)
    structure = handler.detect_structure()
    data = handler.read_data(structure)
    for item in data:
        item["adjusted_price"] = item["price"]

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "test_output.xlsx")
        handler.write_results(data, structure, output_path)
        assert os.path.exists(output_path)

        wb = load_workbook(output_path)
        ws = wb[structure["sheet_name"]]
        header_cell = ws.cell(row=structure["header_row"], column=structure["col_price"] + 1)
        assert header_cell.value == "调价后合价"
        wb.close()


def test_invalid_input_no_price_column():
    with tempfile.TemporaryDirectory() as tmpdir:
        wb = Workbook()
        ws = wb.active
        ws.title = "测试表"
        ws.append(["序号", "项目编码", "项目名称"])
        ws.append([1, "010101001", "测试项目"])
        filepath = os.path.join(tmpdir, "no_price.xlsx")
        wb.save(filepath)
        wb.close()

        handler = ExcelHandler(filepath)
        with pytest.raises(ValueError, match="ERR_INVALID_INPUT"):
            handler.detect_structure()
