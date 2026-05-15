import gradio as gr
import pandas as pd
import numpy as np
from lite_processor import LiteDataProcessor

processor = LiteDataProcessor()

INPUT_HEADERS = ["项目名称", "工程量", "综合单价", "合价", "单价浮动上限%", "单价浮动下限%", "是否锁定"]
RESULT_HEADERS = ["项目名称", "工程量", "原综合单价", "原合价", "调价后综合单价", "调价后合价", "浮动比例%", "是否锁定"]

EMPTY_ROW = ["", None, None, None, None, None, False]


def make_empty_df(n=10):
    return pd.DataFrame([EMPTY_ROW[:] for _ in range(n)], columns=INPUT_HEADERS)


def on_auto_calc(input_df):
    if input_df is None or input_df.empty:
        return input_df
    df = input_df.copy()
    for i in range(len(df)):
        try:
            qty = df.at[i, "工程量"]
            up = df.at[i, "综合单价"]
            if qty is not None and up is not None and qty != "" and up != "":
                qty_f = float(qty)
                up_f = float(up)
                df.at[i, "合价"] = round(qty_f * up_f, 2)
            else:
                df.at[i, "合价"] = None
        except (ValueError, TypeError):
            df.at[i, "合价"] = None
    return df


def on_add_rows(input_df):
    if input_df is None:
        df = make_empty_df(5)
    else:
        df = input_df.copy()
        for _ in range(5):
            df = pd.concat([df, pd.DataFrame([EMPTY_ROW[:]], columns=INPUT_HEADERS)], ignore_index=True)
    return df


def on_clear():
    return make_empty_df(10)


def df_to_gui_rows(input_df):
    gui_rows = []
    if input_df is None:
        return gui_rows
    for i in range(len(input_df)):
        row = input_df.iloc[i]
        name = str(row["项目名称"]) if row["项目名称"] is not None and str(row["项目名称"]).strip() != "" else ""
        if name == "nan":
            name = ""

        try:
            quantity = float(row["工程量"]) if row["工程量"] is not None and str(row["工程量"]).strip() != "" else 0
        except (ValueError, TypeError):
            quantity = 0

        try:
            unit_price = float(row["综合单价"]) if row["综合单价"] is not None and str(row["综合单价"]).strip() != "" else 0
        except (ValueError, TypeError):
            unit_price = 0

        price = round(quantity * unit_price, 2)

        try:
            float_up = float(row["单价浮动上限%"]) if row["单价浮动上限%"] is not None and str(row["单价浮动上限%"]).strip() != "" else None
        except (ValueError, TypeError):
            float_up = None

        try:
            float_down = float(row["单价浮动下限%"]) if row["单价浮动下限%"] is not None and str(row["单价浮动下限%"]).strip() != "" else None
        except (ValueError, TypeError):
            float_down = None

        is_locked = bool(row["是否锁定"]) if row["是否锁定"] is not None else False

        if name == "" and quantity == 0 and unit_price == 0:
            continue

        gui_rows.append({
            "name": name,
            "quantity": quantity,
            "unit_price": unit_price,
            "price": price,
            "float_up": float_up,
            "float_down": float_down,
            "is_locked": is_locked,
        })
    return gui_rows


def on_calculate(input_df, large_low, large_high, medium_low, medium_high,
                 small_low, small_high, seed, ud_min, ud_max,
                 total_min, total_max):
    gui_rows = df_to_gui_rows(input_df)
    if not gui_rows:
        return make_empty_df(0), "⚠️ 未输入有效数据，请先填写项目信息", ""

    large_range = (large_low / 100.0, large_high / 100.0)
    medium_range = (medium_low / 100.0, medium_high / 100.0)
    small_range = (small_low / 100.0, small_high / 100.0)

    total_price_min = total_min if total_min is not None and total_min > 0 else None
    total_price_max = total_max if total_max is not None and total_max > 0 else None

    up_down_min_val = ud_min / 100.0
    up_down_max_val = ud_max / 100.0

    try:
        result = processor.run_lite_adjustment(
            gui_rows,
            large_range=large_range,
            medium_range=medium_range,
            small_range=small_range,
            total_price_min=total_price_min,
            total_price_max=total_price_max,
            up_down_min=up_down_min_val,
            up_down_max=up_down_max_val,
            seed=int(seed),
        )
    except Exception as e:
        return make_empty_df(0), f"❌ 调价执行出错: {str(e)}", ""

    result_rows = result["result_rows"]
    quality = result["quality"]
    summary = result["summary"]

    rows_data = []
    for r in result_rows:
        rows_data.append([
            r["name"],
            r["quantity"],
            f'{r["unit_price"]:.2f}',
            f'{r["price"]:.2f}',
            f'{r["adjusted_unit_price"]:.2f}',
            f'{r["adjusted_price"]:.2f}',
            f'{r["float_pct"]:.2f}',
            "✅" if r["is_locked"] else "",
        ])

    result_df = pd.DataFrame(rows_data, columns=RESULT_HEADERS)

    summary_text = (
        f"📊 **摘要信息**\n\n"
        f"| 指标 | 值 |\n|---|---|\n"
        f"| 原总价 | {summary['original_total']:,.2f} |\n"
        f"| 新总价 | {summary['adjusted_total']:,.2f} |\n"
        f"| 变化额 | {summary['change_amount']:,.2f} |\n"
        f"| 变化率 | {summary['change_pct']:.2f}% |\n"
        f"| 锁定项 | {summary['locked_count']} |\n"
        f"| 可调项 | {summary['adjustable_count']} |\n"
        f"| 上涨项 | {summary['up_count']} |\n"
        f"| 下跌项 | {summary['down_count']} |"
    )

    quality_text = format_quality_report(quality, summary)

    return result_df, summary_text, quality_text


def format_quality_report(quality, summary):
    lines = []
    lines.append("## 🔍 质量检验报告\n")

    imbalance = quality.get("imbalance_detection", {})
    direction = quality.get("direction_detection", {})
    correlation = quality.get("correlation_detection", {})
    tier = quality.get("tier_compliance", {})

    lines.append("### 1. 不平衡检测\n")
    gini = imbalance.get("gini_coefficient", {})
    gini_status = "✅ 通过" if gini.get("passed") else "❌ 未通过"
    lines.append(f"- **基尼系数**: {gini.get('value', 'N/A')} (标准: {gini.get('standard', 'N/A')}) → {gini_status}")

    z_outlier = imbalance.get("z_score_outliers", {})
    z_val = z_outlier.get("value", {})
    z_status = "✅ 通过" if z_outlier.get("passed") else "❌ 未通过"
    lines.append(f"- **Z-score异常值**: {z_val.get('count', 0)}个 ({z_val.get('pct', 0)}%) (标准: {z_outlier.get('standard', 'N/A')}) → {z_status}")

    iqr_outlier = imbalance.get("iqr_outliers", {})
    iqr_status = "✅ 通过" if iqr_outlier.get("passed") else "❌ 未通过"
    lines.append(f"- **IQR异常值**: {iqr_outlier.get('value', 0)}个 (标准: {iqr_outlier.get('standard', 'N/A')}) → {iqr_status}")

    lines.append("\n### 2. 涨跌方向检测\n")
    ud_ratio = direction.get("up_down_ratio", {})
    ud_status = "✅ 通过" if ud_ratio.get("passed") else "❌ 未通过"
    lines.append(f"- **涨跌比例**: {ud_ratio.get('value', 'N/A')} (标准: {ud_ratio.get('standard', 'N/A')}) → {ud_status}")

    binom = direction.get("binomial_test_p", {})
    binom_status = "✅ 通过" if binom.get("passed") else "❌ 未通过"
    lines.append(f"- **二项检验p值**: {binom.get('value', 'N/A')} (标准: {binom.get('standard', 'N/A')}) → {binom_status}")

    lines.append("\n### 3. 相关性检测\n")
    pearson = correlation.get("pearson_r", {})
    pearson_status = "✅ 通过" if pearson.get("passed") else "❌ 未通过"
    lines.append(f"- **Pearson r**: {pearson.get('value', 'N/A')} (标准: {pearson.get('standard', 'N/A')}) → {pearson_status}")

    spearman = correlation.get("spearman_rho", {})
    spearman_status = "✅ 通过" if spearman.get("passed") else "❌ 未通过"
    lines.append(f"- **Spearman ρ**: {spearman.get('value', 'N/A')} (标准: {spearman.get('standard', 'N/A')}) → {spearman_status}")

    ks = correlation.get("ks_test_p", {})
    ks_status = "✅ 通过" if ks.get("passed") else "❌ 未通过"
    lines.append(f"- **KS检验p值**: {ks.get('value', 'N/A')} (标准: {ks.get('standard', 'N/A')}) → {ks_status}")

    lines.append("\n### 4. 层级合规性\n")
    lmf = tier.get("large_max_float", {})
    lmf_status = "✅ 通过" if lmf.get("passed") else "❌ 未通过"
    lines.append(f"- **大项最大浮动**: {lmf.get('value', 'N/A')}% (标准: {lmf.get('standard', 'N/A')}) → {lmf_status}")

    smf = tier.get("small_min_float", {})
    smf_status = "✅ 通过" if smf.get("passed") else "❌ 未通过"
    lines.append(f"- **小项最小浮动**: {smf.get('value', 'N/A')}% (标准: {smf.get('standard', 'N/A')}) → {smf_status}")

    fps = tier.get("float_price_spearman", {})
    fps_status = "✅ 通过" if fps.get("passed") else "❌ 未通过"
    lines.append(f"- **浮动-价格Spearman**: {fps.get('value', 'N/A')} (标准: {fps.get('standard', 'N/A')}) → {fps_status}")

    lines.append(f"\n### 5. 总价变化率\n")
    lines.append(f"- **变化率**: {summary.get('change_pct', 0):.2f}%")

    fail_count = 0
    warn_count = 0
    for section in [imbalance, direction, correlation, tier]:
        for key, check in section.items():
            if isinstance(check, dict) and "passed" in check:
                if not check["passed"]:
                    fail_count += 1

    lines.append(f"\n### 6. 结论\n")
    if fail_count == 0:
        lines.append("✅ **所有检验项均通过**，调价结果质量良好。")
    elif fail_count <= 2:
        lines.append(f"⚠️ **{fail_count}项检验未通过**，建议关注并调整参数。")
    else:
        lines.append(f"❌ **{fail_count}项检验未通过**，建议重新调整参数或更换随机种子。")

    return "\n".join(lines)


def on_export(result_df):
    if result_df is None or result_df.empty:
        return None

    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "调价结果"

    headers = ["项目名称", "工程量", "原综合单价", "原合价", "调价后综合单价", "调价后合价", "浮动比例%", "是否锁定"]
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx in range(len(result_df)):
        for col_idx in range(len(result_df.columns)):
            val = result_df.iloc[row_idx, col_idx]
            if isinstance(val, str) and val == "":
                val = ""
            cell = ws.cell(row=row_idx + 2, column=col_idx + 1, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [20, 12, 15, 15, 15, 15, 12, 10]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    filepath = "/tmp/调价结果.xlsx"
    wb.save(filepath)
    return filepath


with gr.Blocks(title="工程造价智能调价工具 Lite") as demo:
    gr.Markdown("# 工程造价智能调价工具 Lite")

    with gr.Accordion("数据输入", open=True):
        input_df = gr.Dataframe(
            headers=INPUT_HEADERS,
            datatype=["str", "number", "number", "number", "number", "number", "bool"],
            row_count=10,
            column_count=len(INPUT_HEADERS),
            interactive=True,
            label="项目数据",
        )
        with gr.Row():
            add_btn = gr.Button("添加5行")
            clear_btn = gr.Button("清空")

    with gr.Accordion("算法参数", open=True):
        with gr.Row():
            large_low = gr.Number(label="大项浮动下限%", value=0.5, precision=2)
            large_high = gr.Number(label="大项浮动上限%", value=1.0, precision=2)
        with gr.Row():
            medium_low = gr.Number(label="中项浮动下限%", value=1.0, precision=2)
            medium_high = gr.Number(label="中项浮动上限%", value=2.0, precision=2)
        with gr.Row():
            small_low = gr.Number(label="小项浮动下限%", value=2.0, precision=2)
            small_high = gr.Number(label="小项浮动上限%", value=3.0, precision=2)
        with gr.Row():
            seed_input = gr.Number(label="随机种子", value=42, precision=0)
        with gr.Row():
            ud_min = gr.Number(label="涨跌比例下限%", value=40, precision=1)
            ud_max = gr.Number(label="涨跌比例上限%", value=60, precision=1)
        with gr.Row():
            total_min = gr.Number(label="总价下限", value=None, precision=2)
            total_max = gr.Number(label="总价上限", value=None, precision=2)

    calc_btn = gr.Button("🔧 执行调价", variant="primary")

    with gr.Accordion("调价结果", open=True):
        summary_md = gr.Markdown("")
        result_df = gr.Dataframe(
            headers=RESULT_HEADERS,
            datatype=["str", "number", "str", "str", "str", "str", "str", "str"],
            interactive=False,
            label="调价结果",
        )
        export_btn = gr.Button("📥 导出Excel")
        export_file = gr.File(label="下载文件")

    with gr.Accordion("质量检验报告", open=True):
        quality_md = gr.Markdown("")

    input_df.change(fn=on_auto_calc, inputs=[input_df], outputs=[input_df])

    add_btn.click(fn=on_add_rows, inputs=[input_df], outputs=[input_df])
    clear_btn.click(fn=on_clear, inputs=None, outputs=[input_df])

    calc_btn.click(
        fn=on_calculate,
        inputs=[input_df, large_low, large_high, medium_low, medium_high,
                small_low, small_high, seed_input, ud_min, ud_max,
                total_min, total_max],
        outputs=[result_df, summary_md, quality_md],
    )

    export_btn.click(fn=on_export, inputs=[result_df], outputs=[export_file])

if __name__ == "__main__":
    demo.launch(server_name="0.0.0.0", server_port=7860)
