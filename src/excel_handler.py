import shutil
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class ExcelHandler:
    def __init__(self, filepath):
        self.filepath = filepath
        self.wb = load_workbook(filepath, data_only=True)

    def detect_structure(self):
        code_keywords = ["编码", "项目编码", "编号"]
        name_keywords = ["名称", "项目名称"]
        price_keywords = ["合价", "金额"]

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=False), start=1):
                cell_values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
                col_code = None
                col_name = None
                col_price = None
                for col_idx, val in enumerate(cell_values, start=1):
                    if any(kw in val for kw in code_keywords) and col_code is None:
                        col_code = col_idx
                    if any(kw in val for kw in name_keywords) and col_name is None:
                        col_name = col_idx
                    if any(kw in val for kw in price_keywords) and col_price is None:
                        col_price = col_idx
                if col_price is not None:
                    data_start_row = row_idx + 1
                    total_rows = 0
                    for r in ws.iter_rows(min_row=data_start_row, max_col=1, values_only=True):
                        if r[0] is not None:
                            total_rows += 1
                        else:
                            break
                    return {
                        "sheet_name": sheet_name,
                        "header_row": row_idx,
                        "data_start_row": data_start_row,
                        "col_code": col_code,
                        "col_name": col_name,
                        "col_price": col_price,
                        "total_rows": total_rows,
                    }

        raise ValueError("ERR_INVALID_INPUT: 无法找到合价列")

    def read_data(self, structure):
        ws = self.wb[structure["sheet_name"]]
        result = []
        for row_idx in range(structure["data_start_row"], structure["data_start_row"] + structure["total_rows"]):
            code = ws.cell(row=row_idx, column=structure["col_code"]).value
            name = ws.cell(row=row_idx, column=structure["col_name"]).value
            price = ws.cell(row=row_idx, column=structure["col_price"]).value
            if price is not None:
                try:
                    price = float(price)
                except (ValueError, TypeError):
                    price = None
            result.append({
                "row": row_idx,
                "code": str(code) if code is not None else "",
                "name": str(name) if name is not None else "",
                "price": price,
            })
        return result

    def write_results(self, data, structure, output_path):
        shutil.copy2(self.filepath, output_path)
        wb = load_workbook(output_path)
        ws = wb[structure["sheet_name"]]

        new_col = structure["col_price"] + 1
        header_row = structure["header_row"]
        ws.cell(row=header_row, column=new_col, value="调价后合价")

        src_cell = ws.cell(row=header_row, column=structure["col_price"])
        dst_header_cell = ws.cell(row=header_row, column=new_col)
        if src_cell.font:
            dst_header_cell.font = src_cell.font.copy()
        if src_cell.border:
            dst_header_cell.border = src_cell.border.copy()
        if src_cell.alignment:
            dst_header_cell.alignment = src_cell.alignment.copy()
        if src_cell.number_format:
            dst_header_cell.number_format = src_cell.number_format

        for item in data:
            row_idx = item["row"]
            price_cell = ws.cell(row=row_idx, column=structure["col_price"])
            new_cell = ws.cell(row=row_idx, column=new_col)
            new_cell.value = item.get("adjusted_price", item.get("price"))

            if price_cell.font:
                new_cell.font = price_cell.font.copy()
            if price_cell.border:
                new_cell.border = price_cell.border.copy()
            if price_cell.alignment:
                new_cell.alignment = price_cell.alignment.copy()
            if price_cell.number_format:
                new_cell.number_format = price_cell.number_format

        wb.save(output_path)
        wb.close()
