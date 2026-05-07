import os
import glob
import platform
import numpy as np
import pandas as pd
from openpyxl import load_workbook


class BatchProcessor:
    def __init__(self, config=None):
        self.config = config
        self._results = []
        self._success_count = 0
        self._failed_files = []

    def process_directory(self, input_dir, output_dir, seed=42, discount=None, tolerance=0.01):
        os.makedirs(output_dir, exist_ok=True)
        xlsx_files = sorted(glob.glob(os.path.join(input_dir, "*.xlsx")))
        return self._process_files(xlsx_files, output_dir, seed, discount, tolerance)

    def process_file_list(self, file_list, output_dir, seed=42, discount=None, tolerance=0.01):
        os.makedirs(output_dir, exist_ok=True)
        return self._process_files(file_list, output_dir, seed, discount, tolerance)

    def _process_files(self, file_list, output_dir, seed, discount, tolerance):
        self._results = []
        self._success_count = 0
        self._failed_files = []
        total = len(file_list)

        for idx, input_path in enumerate(file_list, start=1):
            filename = os.path.basename(input_path)
            output_path = os.path.join(output_dir, filename)
            print(f"[{idx}/{total}] 处理: {filename}")

            try:
                result = self._process_single_file(input_path, output_path, seed, discount, tolerance)
                self._results.append(result)
                self._success_count += 1
            except Exception as e:
                self._failed_files.append({"file": input_path, "error": str(e)})
                self._results.append({"file": input_path, "status": "failed", "error": str(e)})
                print(f"  错误: {e}")

        return self.get_summary()

    def _process_single_file(self, input_path, output_path, seed, discount, tolerance):
        from src.excel_handler import ExcelHandler
        from src.data_processor import DataProcessor
        from src.report_generator import ReportGenerator

        handler = ExcelHandler(input_path)
        structure = handler.detect_structure()
        data = handler.read_data(structure)

        processor = DataProcessor(seed=seed)
        data = processor.identify_locked_items(data)
        data = processor.tier_grouping(data)
        data = processor.generate_coefficients(data)
        data = processor.control_up_down_ratio(data)
        data = processor.control_correlation(data)
        data = processor.calibrate_total_price(data, target_discount_pct=discount, tolerance=tolerance)

        handler.write_results(data, structure, output_path)

        quality = processor.advanced_quality_check(data)

        report_config = self.config if self.config else None
        if report_config is None:
            report_config = type("Config", (), {"get": lambda self, k: {"input_file": input_path, "random_seed": seed}.get(k)})()
        else:
            if self.config.get("input_file") is None:
                self.config._config["input_file"] = input_path
            if self.config.get("random_seed") is None:
                self.config._config["random_seed"] = seed

        generator = ReportGenerator()
        report_text = generator.generate_report(data, quality, config=report_config)

        report_path = output_path.replace(".xlsx", "_report.txt")
        generator.save_report(report_text, report_path)

        original_total = sum(item.get("price", 0) or 0 for item in data)
        adjusted_total = sum(item.get("adjusted_price", 0) or 0 for item in data)
        change_pct = round((adjusted_total - original_total) / original_total * 100, 2) if original_total else 0

        return {
            "file": input_path,
            "status": "success",
            "quality": quality,
            "original_total": original_total,
            "adjusted_total": adjusted_total,
            "change_pct": change_pct,
        }

    def get_summary(self):
        return {
            "total_files": len(self._results),
            "success_files": self._success_count,
            "failed_files": self._failed_files,
            "results": self._results,
        }

    def _should_use_openpyxl_acceleration(self, filepath):
        if filepath.endswith(".xlsm"):
            return False

        try:
            import zipfile
            with zipfile.ZipFile(filepath, 'r') as zf:
                if 'vbaProject.bin' in zf.namelist():
                    return False
        except Exception:
            pass

        try:
            wb = load_workbook(filepath, data_only=False)
            ws = wb[wb.sheetnames[0]]
            row_count = 0
            has_formula = False
            for row in ws.iter_rows(min_row=1, values_only=True):
                row_count += 1
                if row_count > 100:
                    break
                for cell_val in row:
                    if cell_val is not None and isinstance(cell_val, str) and cell_val.startswith("="):
                        has_formula = True
                        break
                if has_formula:
                    break
            wb.close()

            if has_formula:
                return False

            if row_count <= 10000:
                return False

            return True
        except Exception:
            return False

    def _accelerated_read(self, filepath, sheet_name=None):
        df = pd.read_excel(filepath, sheet_name=sheet_name or 0, engine='openpyxl', dtype={'编号': str})
        return df

    def _accelerated_write(self, df, filepath):
        df.to_excel(filepath, index=False, engine='openpyxl')

        try:
            wb = load_workbook(filepath, data_only=False)
            has_formula = False
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cell_val in row:
                        if cell_val is not None and isinstance(cell_val, str) and cell_val.startswith("="):
                            has_formula = True
                            break
                    if has_formula:
                        break
                if has_formula:
                    break
            wb.close()

            if has_formula:
                if platform.system() == "Linux":
                    print("WARNING: 文件包含公式，Linux环境下无法使用xlwings重算，请手动打开Excel验证")
                else:
                    try:
                        import xlwings as xw
                        app = xw.App(visible=False)
                        wb_xw = app.books.open(filepath)
                        wb_xw.calculate()
                        wb_xw.save()
                        wb_xw.close()
                        app.quit()
                    except Exception as e:
                        print(f"WARNING: xlwings重算失败: {e}")
        except Exception:
            pass
